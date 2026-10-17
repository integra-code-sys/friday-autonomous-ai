#!/usr/bin/env python3
"""ftools — ferramentas determinísticas da F.R.I.D.A.Y. em CLI (separadas do LLM para errar menos).

  ftools.py extrato ARQ [--saida X.xlsx]   CSV/OFX/PDF de banco -> planilha (Lançamentos, Categorias, Mensal) + resumo JSON
  ftools.py ler ARQ                        despeja xlsx/csv/docx como texto
  ftools.py sudo "comando"                 roda como root via sudo -A (a senha é digitada pelo dono numa janela zenity; nunca guardada)

Saída sempre JSON numa linha: {"ok":bool, ...}. Código de saída 0 = ok.
"""
import csv, io, json, os, re, shlex, subprocess, sys, tempfile, time
from collections import defaultdict
from datetime import datetime

SAIDA_DIR = os.path.expanduser("~/Documentos/Friday")

CATEGORIAS = [   # (categoria, regex no histórico) — primeira que casar vence
    ("Salário/Renda", r"sal[áa]rio|pagto? sal|proventos|remunera|pr[óo]-?labore"),
    ("Transferência", r"\bpix\b|\bted\b|\bdoc\b|transf"),
    ("Mercado", r"mercad|supermerc|atacad|assa[íi]|carrefour|extra |p[ãa]o de a[çc]|hortifr|a[çc]ougue|padaria"),
    ("Alimentação", r"ifood|rappi|restaur|lanch|pizz|burguer|burger|mc ?donald|bar |cafe|caf[ée]"),
    ("Transporte", r"uber|99 ?app|99pop|taxi|posto|combust|gasolin|shell|ipiranga|petrobr|estacion|ped[áa]gio|sem parar"),
    ("Saúde", r"farm[áa]c|drog|raia|pague menos|hospital|cl[íi]nica|laborat|unimed|odonto|m[ée]dic"),
    ("Moradia", r"aluguel|condom|energia|enel|cemig|copel|light|celesc|sabesp|[áa]gua|g[áa]s|iptu"),
    ("Internet/Telefone", r"vivo|claro|\btim\b|\boi\b|net |internet|telecom|fibra"),
    ("Assinaturas", r"netflix|spotify|amazon prime|prime video|disney|hbo|max |youtube|apple\.com|google|microsoft|adobe|anthropic|openai|chatgpt"),
    ("Compras", r"mercado ?livre|mercadopago|amazon|shopee|magalu|magazine|americanas|aliexpress|shein|loja"),
    ("Tarifas/Juros", r"tarifa|juros|iof|anuidade|encargo|multa|cesta"),
    ("Investimentos", r"aplica[çc]|resgate|cdb|tesouro|poupan|invest|corretora"),
    ("Impostos", r"darf|das |simples nac|inss|irpf|detran|ipva|licenciamento"),
    ("Saque", r"saque"),
]


def out(d, code=0):
    print(json.dumps(d, ensure_ascii=False)); sys.exit(code)


def categoria(hist):
    h = hist.lower()
    return next((c for c, rx in CATEGORIAS if re.search(rx, h)), "Outros")


def num(s):
    """'1.234,56' | '-50,00' | '1234.56' | '50,00 D' -> float"""
    s = str(s).strip().replace("R$", "").replace(" ", "")
    neg = s.startswith("-") or s.endswith("-") or s.upper().endswith("D") or (s.startswith("(") and s.endswith(")"))
    s = re.sub(r"[^\d,.]", "", s)
    if not s:
        raise ValueError("vazio")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(".", "").replace(",", ".")
    v = float(s)
    return -v if neg else v


def data(s, ano=None):
    s = str(s).strip()[:10]
    for f in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%Y%m%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    m = re.match(r"(\d{2})/(\d{2})$", s)
    if m:
        return datetime(ano or datetime.now().year, int(m.group(2)), int(m.group(1))).date()
    raise ValueError("data: " + s)


def ler_csv(p):
    raw = open(p, "rb").read()
    txt = next((raw.decode(e) for e in ("utf-8-sig", "latin-1") if _dec(raw, e)), raw.decode("latin-1"))
    linhas = [l for l in txt.splitlines() if l.strip()]
    try:
        dialeto = csv.Sniffer().sniff("\n".join(linhas[:20]), delimiters=";,\t|")
    except csv.Error:
        dialeto = csv.excel; dialeto.delimiter = ";"
    rows = list(csv.reader(linhas, dialeto))
    # cabeçalho = 1ª linha com algo parecido com "data"
    hi = next((i for i, r in enumerate(rows[:15]) if any(re.search(r"\bdata\b|\bdate\b", c, re.I) for c in r)), None)
    if hi is None:
        raise ValueError("CSV sem cabeçalho com coluna 'data'")
    cab = [c.strip().lower() for c in rows[hi]]
    col = lambda *rx: next((i for i, c in enumerate(cab) if any(re.search(r, c) for r in rx)), None)
    cd = col(r"^data|date")
    ch = col(r"hist|descri|lan[çc]amento|memo|estabelec|title|detalhe")
    cv = col(r"^valor|amount|^value|quantia")
    cc, cdb = col(r"cr[ée]dito|entrada"), col(r"d[ée]bito|sa[íi]da")
    if cd is None or (cv is None and cc is None and cdb is None):
        raise ValueError(f"colunas não reconhecidas: {cab}")
    lanc = []
    for r in rows[hi + 1:]:
        try:
            d = data(r[cd])
            if cv is not None and r[cv].strip():
                v = num(r[cv])
            else:
                v = (num(r[cc]) if cc is not None and r[cc].strip() else 0) - (abs(num(r[cdb])) if cdb is not None and r[cdb].strip() else 0)
            h = " ".join(r[ch].split()) if ch is not None else ""
        except (ValueError, IndexError):
            continue
        if re.search(r"saldo", h, re.I):
            continue
        lanc.append((d, h, v))
    return lanc


def _dec(raw, enc):
    try:
        raw.decode(enc); return True
    except UnicodeDecodeError:
        return False


def ler_ofx(p):
    txt = open(p, "rb").read().decode("latin-1")
    lanc = []
    for b in re.findall(r"<STMTTRN>(.*?)(?:</STMTTRN>|(?=<STMTTRN>)|</BANKTRANLIST>)", txt, re.S | re.I):
        g = lambda t: (re.search(rf"<{t}>([^<\r\n]*)", b, re.I) or [None, ""])[1].strip()
        try:
            lanc.append((data(g("DTPOSTED")[:8]), g("MEMO") or g("NAME"), float(g("TRNAMT").replace(",", "."))))
        except ValueError:
            continue
    return lanc


def ler_pdf(p):
    txt = subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, text=True, timeout=60).stdout
    ano = re.search(r"\b(20\d{2})\b", txt)
    ano = int(ano.group(1)) if ano else None
    rx = re.compile(r"^\s*(\d{2}/\d{2}(?:/\d{2,4})?)\s+(.+?)\s{2,}(-?\(?[\d.]+,\d{2}\)?\s?[DC-]?)(?:\s|$)")
    lanc = []
    for l in txt.splitlines():
        m = rx.match(l)
        if not m or re.search(r"saldo", m.group(2), re.I):
            continue
        try:
            v = num(m.group(3))
            if m.group(3).strip().upper().endswith("C"):
                v = abs(v)
            lanc.append((data(m.group(1), ano), " ".join(m.group(2).split()), v))
        except ValueError:
            continue
    return lanc


def cmd_extrato(args):
    if not args:
        out({"ok": False, "erro": "uso: extrato ARQ [--saida X.xlsx]"}, 2)
    p = os.path.realpath(os.path.expanduser(args[0]))
    if not os.path.isfile(p):
        out({"ok": False, "erro": "não existe: " + p}, 1)
    ext = p.lower().rsplit(".", 1)[-1]
    try:
        lanc = {"ofx": ler_ofx, "pdf": ler_pdf}.get(ext, ler_csv)(p)
    except Exception as e:
        out({"ok": False, "erro": str(e)[:200]}, 1)
    if not lanc:
        out({"ok": False, "erro": "nenhum lançamento reconhecido (formato do banco não suportado?)"}, 1)
    lanc.sort(key=lambda x: x[0])
    saida = args[args.index("--saida") + 1] if "--saida" in args else os.path.join(
        SAIDA_DIR, f"Analise extrato {os.path.splitext(os.path.basename(p))[0]} {time.strftime('%Y-%m-%d %H%M')}.xlsx")
    saida = os.path.expanduser(saida); os.makedirs(os.path.dirname(saida), exist_ok=True)

    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
    neg = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="FF8A1F")
    wb = openpyxl.Workbook()

    def aba(ws, cab, linhas, moeda_cols, larg):
        ws.append(cab)
        for c in ws[1]:
            c.font, c.fill = neg, fill
        for l in linhas:
            ws.append(l)
        for col in moeda_cols:
            for c in ws[get_column_letter(col)][1:]:
                c.number_format = BRL
        for i, w in enumerate(larg, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

    ws = wb.active; ws.title = "Lançamentos"
    aba(ws, ["Data", "Histórico", "Categoria", "Valor"], [(d, h, categoria(h), v) for d, h, v in lanc], [4], [12, 60, 20, 16])
    for c in ws["A"][1:]:
        c.number_format = "DD/MM/YYYY"
    n = len(lanc) + 1
    cats = sorted({categoria(h) for _, h, _ in lanc})
    aba(wb.create_sheet("Categorias"), ["Categoria", "Entradas", "Saídas", "Saldo", "Qtd"],
        [(c, f'=SUMIFS(Lançamentos!D2:D{n},Lançamentos!C2:C{n},"{c}",Lançamentos!D2:D{n},">0")',
          f'=SUMIFS(Lançamentos!D2:D{n},Lançamentos!C2:C{n},"{c}",Lançamentos!D2:D{n},"<0")',
          f"=B{i}+C{i}", f'=COUNTIF(Lançamentos!C2:C{n},"{c}")') for i, c in enumerate(cats, 2)], [2, 3, 4], [22, 16, 16, 16, 8])
    mes = defaultdict(lambda: [0.0, 0.0])
    for d, _, v in lanc:
        mes[d.strftime("%Y-%m")][0 if v > 0 else 1] += v
    aba(wb.create_sheet("Mensal"), ["Mês", "Entradas", "Saídas", "Saldo"],
        [(m, round(e, 2), round(s, 2), f"=B{i}+C{i}") for i, (m, (e, s)) in enumerate(sorted(mes.items()), 2)], [2, 3, 4], [10, 16, 16, 16])
    wb.save(saida)

    ent = sum(v for *_, v in lanc if v > 0); sai = sum(v for *_, v in lanc if v < 0)
    por_cat = defaultdict(float)
    for _, h, v in lanc:
        if v < 0:
            por_cat[categoria(h)] += v
    top = sorted(por_cat.items(), key=lambda x: x[1])[:5]
    maiores = sorted((x for x in lanc if x[2] < 0), key=lambda x: x[2])[:5]
    out({"ok": True, "arquivo": saida, "lancamentos": len(lanc), "periodo": [str(lanc[0][0]), str(lanc[-1][0])],
         "entradas": round(ent, 2), "saidas": round(sai, 2), "saldo": round(ent + sai, 2),
         "top_categorias_gasto": [[c, round(v, 2)] for c, v in top],
         "maiores_gastos": [[str(d), h[:50], v] for d, h, v in maiores],
         "outros_sem_categoria": sum(1 for _, h, v in lanc if categoria(h) == "Outros")})


def cmd_ler(args):
    p = os.path.realpath(os.path.expanduser(args[0])) if args else ""
    if not os.path.isfile(p):
        out({"ok": False, "erro": "não existe: " + p}, 1)
    ext = p.lower().rsplit(".", 1)[-1]
    if ext == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        txt = "\n".join(f"## {ws.title}\n" + "\n".join(" | ".join("" if c is None else str(c) for c in r)
                                                      for r in ws.iter_rows(values_only=True, max_row=300)) for ws in wb.worksheets)
    elif ext == "docx":
        import docx
        d = docx.Document(p)
        txt = "\n".join(x.text for x in d.paragraphs) + "".join(
            "\n" + "\n".join(" | ".join(c.text for c in r.cells) for r in t.rows) for t in d.tables)
    else:
        txt = open(p, "rb").read().decode("utf-8", "replace")
    out({"ok": True, "texto": txt[:20000]})


PERIGOSO = re.compile(r"rm\s+-[a-z]*r[a-z]*f?\s+/(\s|$|\*)|\bmkfs|\bdd\b.*\bof=/dev/|>\s*/dev/sd|:\(\)\s*\{|chmod\s+-R\s+[0-7]+\s+/(\s|$)|"
                      r"\bshred\b|wipefs|fdisk|parted|cryptsetup|passwd\b|visudo|/etc/sudoers|usermod|userdel", re.I)


def cmd_sudo(args):
    c = " ".join(args).strip()
    if not c:
        out({"ok": False, "erro": "comando vazio"}, 2)
    if PERIGOSO.search(c):
        out({"ok": False, "erro": "comando bloqueado por segurança (destrutivo/contas/discos): faça manualmente"}, 3)
    import shell_seguro                              # F7: mesma classificação do terminal (NEGADO não roda) e roda num scope com limite
    plano = shell_seguro.rodar_root(c, 900)
    if plano.get("ok") is False:
        out({"ok": False, "erro": plano["erro"]}, 3)
    ask = os.path.join(tempfile.gettempdir(), "friday-askpass.sh")
    with open(ask, "w") as f:
        f.write('#!/bin/sh\nexec zenity --password --title="F.R.I.D.A.Y. — sudo" 2>/dev/null\n')
    os.chmod(ask, 0o700)
    env = dict(os.environ, SUDO_ASKPASS=ask)
    try:
        r = subprocess.run(plano["cmd"], capture_output=True, text=True, timeout=960, env=env)
    except subprocess.TimeoutExpired:
        out({"ok": False, "erro": "tempo esgotado (15 min)"}, 1)
    try:
        d = json.loads(r.stdout.strip().splitlines()[-1])           # saída do executor do shell_seguro (já cortada em 3 KB)
    except Exception:
        out({"ok": False, "codigo": r.returncode, "saida": ("senha recusada/cancelada ou tempo esgotado: " + (r.stderr or r.stdout)).strip()[-600:]}, 1)
    out({"ok": d["codigo"] == 0, "codigo": d["codigo"], "saida": d["saida"]}, 0 if d["codigo"] == 0 else 1)


if __name__ == "__main__":
    if len(sys.argv) < 2 or sys.argv[1] not in ("extrato", "ler", "sudo"):
        print(__doc__); sys.exit(2)
    {"extrato": cmd_extrato, "ler": cmd_ler, "sudo": cmd_sudo}[sys.argv[1]](sys.argv[2:])
